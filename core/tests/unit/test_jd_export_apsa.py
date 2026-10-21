"""Track J, PR J4 — the APSA/APEX/Poly export reads back, and looks like HR's template.

Two properties, both measured against the Bank's OWN ingest path (``extract_text`` ->
``parse_jd``), never against the renderer's own idea of what it wrote:

1. **Round trip.** Before J4, 1,798 of 1,876 live non-CUPE roles re-read with a
   different duty count (usually "5 exported, 1 re-read"): the export printed duties as
   Word "List Bullet" paragraphs, whose bullet is list FORMATTING with no character in
   the extracted text, so the JDFN parser joined every duty — and the footer, under a
   heading the parser does not know — into one. J4 fixes it on the RENDERER side: each
   item is printed with a typed bullet the parser reads (``apsa_form.bullet_glyph``,
   HR-259), how/why clauses as Word sub-bullets under their duty, and every heading in
   the parser's own vocabulary (``CANONICAL_HEADING``).
2. **HR's 2024 template** (``APSA_APEX_POLY_JD Template_20240530.docx``): the
   identification table in the page header (Position Title / Position # / Employee
   Group / dates), the About SFU block first, sections in template order.

Plus the GOLDEN file: the export's zip-member content is pinned, so "the APSA export did
not change" is a gate, not a one-off measurement. Regenerate deliberately with
``UPDATE_APSA_GOLDEN=1`` and say in the PR what changed.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from src.jd_bank.db.models import DocumentFormat
from src.jd_bank.ingest.extract import extract_text
from src.jd_core.models.parsed_jd import (
    SFUDuty,
    SFUJobDescription,
    SFUQualification,
    SFURelationships,
)
from src.jd_core.parser import headings as hd
from src.jd_core.parser.headings import CANONICAL_HEADING, SectionKey
from src.jd_core.parser.segmenter import _BULLET_RE, parse_jd
from src.jd_core.rules import Rules, get_rules
from src.jd_export import render_jdfn_docx, render_sfu_docx
from tests.export_roundtrip import jdfn_round_trip_losses

GOLDEN = Path(__file__).resolve().parent.parent / "fixtures" / "golden_export"
GOLDEN_FILE = GOLDEN / "apsa_export_v1.json"


@pytest.fixture(scope="module")
def rules() -> Rules:
    return get_rules()


def _rich_jd(rules: Rules, **update: object) -> SFUJobDescription:
    """An APSA JD exercising every section the export prints — including the J4
    additions (position #, how/why, a security credential, supervisor title / #)."""
    jd = SFUJobDescription(
        title="Academic Programs Manager",
        position_number="00012345",
        employee_group="apsa",
        supervisor_title="Director, Academic Operations",
        supervisor_position_number="00054321",
        about_sfu_present=True,
        territorial_acknowledgement_present=True,
        employment_equity_present=True,
        position_summary=(
            "Provides managerial and operational support for the undergraduate and "
            "graduate programs of the School."
        ),
        duties=[
            SFUDuty(
                action_verb="Manages",
                statement="Manages the delivery of all academic programs by: (40%)",
                how_why=[
                    "Representing the department on relevant committees.",
                    "Developing work plans and evaluating their effectiveness.",
                ],
            ),
            SFUDuty(
                action_verb="Leads",
                statement="Leads staffing and teaching assistant processes (25%)",
            ),
            SFUDuty(
                action_verb="Coordinates",
                statement="Coordinates curriculum planning and marketing (15%)",
                how_why=["Advising faculty on calendar changes."],
            ),
            SFUDuty(
                action_verb="Oversees",
                statement="Oversees the program budget (10%)",
            ),
            SFUDuty(
                action_verb="Resolves",
                statement="Resolves student and staff issues (10%)",
            ),
        ],
        decision_making=[
            "Approves program expenditures within the budget",
            "Determines course scheduling priorities",
            "Recommends staffing levels",
        ],
        problem_solving=[
            "Resolves timetable conflicts independently",
            "Escalates policy exceptions to the Director",
        ],
        relationships=SFURelationships(
            supervisory="Supervises 4 program staff",
            internal=["Faculty members", "Registrar's Office"],
            external=["Accreditation bodies"],
        ),
        qualifications=[
            SFUQualification(
                text=(
                    "Bachelor's degree in Business Administration or a related "
                    "discipline and four years of related experience"
                ),
                kind="education",
            ),
            SFUQualification(
                text="knowledge of University policies",
                kind="knowledge",
                modifier="Excellent",
            ),
            SFUQualification(text="Advanced project management skills", kind="skill"),
            SFUQualification(
                text="Ability to work cooperatively with others", kind="ability"
            ),
            SFUQualification(
                text=rules.apsa_form.security_check_sentence, kind="security"
            ),
        ],
        additional_context="The role is based at the Burnaby campus.",
    )
    return jd.model_copy(update=update)


def _reread(data: bytes) -> tuple[str, SFUJobDescription]:
    body = extract_text(data, DocumentFormat.DOCX)
    return body, parse_jd(body).jd


def _body_texts(data: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(data)).paragraphs if p.text.strip()]


def _header_lines(data: bytes) -> list[str]:
    doc = Document(BytesIO(data))
    lines: list[str] = []
    for paragraph in doc.sections[0].header.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text.strip())
    for table in doc.sections[0].header.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" ".join(c for c in cells if c))
    return lines


# ── 1. The round trip ────────────────────────────────────────────────────────────


def test_an_apsa_export_reads_back_through_the_banks_own_parser(rules: Rules) -> None:
    """Every field the JDFN parser can read comes back: title, position #, group,
    summary, every duty (with its how/why), decisions, problems, contacts,
    qualifications, additional context, and the boilerplate presence flags."""
    jd = _rich_jd(rules)
    _, back = _reread(render_sfu_docx(jd))
    assert jdfn_round_trip_losses(jd, back) == []


def test_the_duty_count_survives_the_round_trip(rules: Rules) -> None:
    """THE pre-J4 defect, asserted on its own so a regression names itself: five
    duties exported came back as one."""
    jd = _rich_jd(rules)
    _, back = _reread(render_sfu_docx(jd))
    assert len(back.duties) == len(jd.duties) == 5


def test_how_why_clauses_stay_with_their_duty(rules: Rules) -> None:
    """The template's "ACTION VERB… WHAT by / How and why" shape: the clauses print as
    sub-bullets under their duty and re-read as PART of that duty, never as duties of
    their own. Since parser v9 they re-read as that duty's ``how_why`` (nested Word list
    paragraphs — owner ruling 2026-10-20), not fused into its statement."""
    jd = _rich_jd(rules)
    _, back = _reread(render_sfu_docx(jd))
    assert len(back.duties) == len(jd.duties)
    for original, reread in zip(jd.duties, back.duties, strict=True):
        assert reread.how_why == [c.strip() for c in original.how_why if c.strip()]
    assert "Representing the department on relevant committees" in " ".join(
        back.duties[0].how_why
    )
    assert "Advising faculty on calendar changes" in " ".join(back.duties[2].how_why)


def test_a_minimal_jd_still_round_trips(rules: Rules) -> None:
    jd = SFUJobDescription(title="Bare Role")
    _, back = _reread(render_sfu_docx(jd))
    assert jdfn_round_trip_losses(jd, back) == []


# ── 2. HR's 2024 template ────────────────────────────────────────────────────────


def test_the_about_sfu_block_is_printed_from_the_rulebook(rules: Rules) -> None:
    """The About SFU block was never exported (gap analysis A4). It is printed from
    ``boilerplate.yaml`` AS IT STANDS — HR-104 is open and the 2024 template carries a
    different paragraph, so the wording is HR's call, not this renderer's."""
    data = render_sfu_docx(_rich_jd(rules))
    texts = _body_texts(data)
    assert texts[0] == CANONICAL_HEADING[SectionKey.ABOUT_SFU]
    for passage in rules.boilerplate.about_sfu:
        assert passage in texts, passage
    _, back = _reread(data)
    assert back.about_sfu_present is True


def test_the_about_sfu_block_is_printed_even_for_a_bare_jd(rules: Rules) -> None:
    """Like the footer: mandated on every APSA JD, so not gated on the draft's flags."""
    texts = _body_texts(render_sfu_docx(SFUJobDescription(title="Bare Role")))
    for passage in rules.boilerplate.about_sfu:
        assert passage in texts


def test_sections_print_in_template_order_with_the_parsers_headings(
    rules: Rules,
) -> None:
    texts = _body_texts(render_sfu_docx(_rich_jd(rules)))
    order = [
        CANONICAL_HEADING[SectionKey.ABOUT_SFU],
        CANONICAL_HEADING[SectionKey.POSITION_SUMMARY],
        CANONICAL_HEADING[SectionKey.DUTIES],
        CANONICAL_HEADING[SectionKey.DECISION_MAKING],
        CANONICAL_HEADING[SectionKey.PROBLEM_SOLVING],
        CANONICAL_HEADING[SectionKey.RELATIONSHIPS],
        CANONICAL_HEADING[SectionKey.QUALIFICATIONS],
        rules.apsa_form.footer_heading,
        CANONICAL_HEADING[SectionKey.ADDITIONAL_CONTEXT],
    ]
    positions = [texts.index(h) for h in order]
    assert positions == sorted(positions), list(zip(order, positions, strict=True))
    # The footer sits between Qualifications and Additional Contextual Information,
    # exactly where the template puts it.
    footer = rules.boilerplate.territorial_acknowledgement[0]
    assert positions[6] < texts.index(footer) < positions[8]


def test_the_identification_table_is_in_the_page_header(rules: Rules) -> None:
    """The template's header: banner + Position Title / Position # / Employee Group /
    dates — plus the supervisor's title and position # the Builder now collects."""
    form = rules.apsa_form
    lines = _header_lines(render_sfu_docx(_rich_jd(rules)))
    for banner in form.banner:
        assert banner in lines
    labels = form.id_labels
    assert f"{labels['title']} Academic Programs Manager" in lines
    assert f"{labels['position_number']} 00012345" in lines
    assert f"{labels['employee_group']} APSA" in lines
    assert f"{labels['supervisor_title']} Director, Academic Operations" in lines
    assert f"{labels['supervisor_position_number']} 00054321" in lines
    assert labels["initial_effective_date"] in lines
    assert labels["latest_revision_date"] in lines


def test_blank_identification_fields_print_their_label_with_no_value(
    rules: Rules,
) -> None:
    """Never guessed: a field the draft does not carry prints as an empty cell."""
    lines = _header_lines(render_sfu_docx(SFUJobDescription(title="Bare Role")))
    labels = rules.apsa_form.id_labels
    assert labels["position_number"] in lines
    assert labels["supervisor_title"] in lines


def test_the_security_credential_prints_under_qualifications(rules: Rules) -> None:
    """Qualifications "must include … E. Security credentials (if applicable)" (S3
    p.9): the criminal-record-check sentence from ``apsa_form.yaml`` (HR-257)."""
    texts = _body_texts(render_sfu_docx(_rich_jd(rules)))
    quals = texts.index(CANONICAL_HEADING[SectionKey.QUALIFICATIONS])
    footer = texts.index(rules.apsa_form.footer_heading)
    sentence = rules.apsa_form.security_check_sentence
    assert any(sentence in t for t in texts[quals:footer])


def test_items_carry_a_bullet_the_parser_can_see(rules: Rules) -> None:
    """The mechanism of the fix (HR-259): every list item is printed with a TYPED
    bullet, because a Word list bullet is formatting and never reaches the text the
    parser reads. How/why clauses are the exception — Word sub-bullets, on purpose, so
    they re-read as part of their duty."""
    doc = Document(BytesIO(render_sfu_docx(_rich_jd(rules))))
    glyph = rules.apsa_form.bullet_glyph
    duty_items = [p.text for p in doc.paragraphs if p.text.startswith(glyph)]
    assert any("Manages the delivery" in t for t in duty_items)
    assert all(_BULLET_RE.match(t) for t in duty_items)
    sub = [p for p in doc.paragraphs if "Representing the department" in p.text]
    assert sub and not sub[0].text.startswith(glyph)
    assert sub[0].style.name == "List Bullet 2"


def test_the_document_still_uses_the_sfu_base_font(rules: Rules) -> None:
    from docx.shared import Pt

    font = Document(BytesIO(render_sfu_docx(_rich_jd(rules)))).styles["Normal"].font
    assert font.name == "Times New Roman"
    assert font.size == Pt(10)


# ── 3. The form's own data is the parser's vocabulary ───────────────────────────


def test_the_modelled_id_labels_are_ones_the_parser_reads(rules: Rules) -> None:
    """The labels for fields the parser reads back must be spellings its own regexes
    match — otherwise the export prints a field the Bank cannot re-read."""
    labels = rules.apsa_form.id_labels
    assert hd.TITLE_LABEL_RX.search(f"{labels['title']} X")
    assert hd.POSITION_NO_LABEL_RX.search(f"{labels['position_number']} 123")
    assert hd.EMPLOYEE_GROUP_LABEL_RX.search(f"{labels['employee_group']} APSA")
    # ...and the supervisor's labels must NOT read as the position's own fields.
    for key in ("supervisor_title", "supervisor_position_number"):
        line = f"{labels[key]} 999"
        assert not hd.TITLE_LABEL_RX.search(line), key
        assert not hd.POSITION_NO_LABEL_RX.search(line), key


def test_the_footer_heading_is_one_the_parser_knows(rules: Rules) -> None:
    heading = rules.apsa_form.footer_heading
    matched = hd.match_heading(heading)
    assert matched is not None and matched[0] is SectionKey.FOOTER


def test_the_bullet_glyph_is_one_the_parser_reads(rules: Rules) -> None:
    assert _BULLET_RE.match(f"{rules.apsa_form.bullet_glyph}\tItem")


# ── 4. The golden file ───────────────────────────────────────────────────────────

#: Members python-docx stamps per save — compared by NOTHING, everything else is.
_UNSTABLE_MEMBERS = frozenset({"docProps/core.xml"})


def _members(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if name not in _UNSTABLE_MEMBERS
        }


def _snapshot(data: bytes) -> dict[str, object]:
    """The golden: a sha256 per zip member (so ANY content change is caught) plus the
    readable text of the body and header (so a failure says what changed)."""
    return {
        "members": _members(data),
        "header": _header_lines(data),
        "body": _body_texts(data),
    }


def test_the_apsa_export_matches_its_golden_file(rules: Rules) -> None:
    """🔴 "APSA unchanged" as a GATE. A change to the APSA export's content — any zip
    member, any paragraph — fails here until the golden is regenerated on purpose
    (``UPDATE_APSA_GOLDEN=1 make gates``) and the PR says what changed."""
    snapshot = _snapshot(render_jdfn_docx(_rich_jd(rules)))
    if os.environ.get("UPDATE_APSA_GOLDEN") == "1":
        GOLDEN.mkdir(parents=True, exist_ok=True)
        GOLDEN_FILE.write_text(
            json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
    golden = json.loads(GOLDEN_FILE.read_text(encoding="utf-8"))
    assert snapshot["header"] == golden["header"]
    assert snapshot["body"] == golden["body"]
    assert snapshot["members"] == golden["members"]


def test_the_golden_is_deterministic(rules: Rules) -> None:
    """Two renders of the same JD agree member for member — the property the golden
    rests on (python-docx's timestamps live only in the excluded core.xml)."""
    jd = _rich_jd(rules)
    assert _members(render_jdfn_docx(jd)) == _members(render_jdfn_docx(jd))


def test_no_raw_bullet_formatting_leaks_into_top_level_items(rules: Rules) -> None:
    """A top-level item printed in a Word list style would carry a SECOND bullet next
    to its typed one; the export prints them in Normal with a hanging indent."""
    doc = Document(BytesIO(render_sfu_docx(_rich_jd(rules))))
    glyph = re.escape(rules.apsa_form.bullet_glyph)
    for p in doc.paragraphs:
        if re.match(glyph, p.text):
            assert not p.style.name.startswith("List"), p.text
