"""Parser v9 — Word list paragraphs are items, and the duty caps are rulebook data.

Three defects, each verified on the live Bank 2026-10-19 and re-measured on a throwaway
copy of it before this change shipped (``docs/plans/PARSER-V9-LISTS-AND-CAPS.md``):

1. **Word list bullets were invisible.** A Word list paragraph carries its bullet as
   FORMATTING (``<w:numPr>``, directly or through its style), not as a glyph in the
   text. python-docx returns the bare words, so the segmenter's ``_logical_lines`` saw
   no bullet marker and appended every item to the previous one: the whole DUTIES
   section became ONE run-on duty, then was cut at 400 characters. Typed-bullet files
   (a ``•`` in the text) never had the problem — the defect was the FORMAT, not the
   content.
2. **Every duty statement was cut at 400 characters** (``_MAX_DUTY_STMT``, mirrored on
   ``SFUDuty.statement``), unregistered.
3. **The WJQ re-read kept 12 functions** (``_MAX_DUTIES``, majors first) — the minor
   functions of a full CUPE questionnaire were dropped.

The fixtures below are shaped like the archive, not like python-docx's defaults: the
modern SFU template writes its duties as ``ListParagraph`` paragraphs with a DIRECT
``numPr`` (measured on ``20260702_00138566_JDFN_APSA_20260626.docx``), and older files
use a numbered STYLE.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches

from src.jd_bank.db.models import DocumentFormat
from src.jd_bank.ingest.extract import LIST_ITEM_MARKER, SUB_ITEM_MARKER, extract_text
from src.jd_core.bank.shingles import shingle_hashes
from src.jd_core.models.parsed_jd import SFUDuty, SFUJobDescription
from src.jd_core.parser import PARSER_VERSION, parse_jd
from src.jd_core.parser.wjq import segment_wjq
from src.jd_core.rules import get_rules

_DUTIES = [
    "Writes strategic funding proposals and develops proposal budgets.",
    "Supports Principal Investigators in writing innovative funding proposals.",
    "Interprets grant regulations and reviews applications for completeness.",
    "Creates presentations, seminars and workshops on funding opportunities.",
    "Maintains familiarity with leading research at SFU.",
    "Liaises with peers at other institutions on current developments.",
]


def _numpr(num_id: str, ilvl: str = "0") -> object:
    return parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/>'
        f'<w:numId w:val="{num_id}"/></w:numPr>'
    )


def _list_para(doc: DocxDocument, text: str, num_id: str = "2") -> None:
    """A ``ListParagraph`` with a DIRECT ``numPr`` — the modern template's shape."""
    p = doc.add_paragraph(text, style="List Paragraph")
    p._p.get_or_add_pPr().append(_numpr(num_id))


def _numbered_style(doc: DocxDocument, style_id: str = "DutyBullet") -> str:
    """A custom paragraph style that carries ``numPr`` itself, based on Normal — the
    older template's shape (numbering inherited, nothing on the paragraph)."""
    styles = doc.styles.element
    styles.append(
        parse_xml(
            f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="{style_id}">'
            f'<w:name w:val="{style_id}"/><w:basedOn w:val="Normal"/>'
            f'<w:pPr><w:numPr><w:numId w:val="7"/></w:numPr></w:pPr></w:style>'
        )
    )
    return style_id


def _blob(doc: DocxDocument) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _jdfn_doc(add_duty: object) -> bytes:
    doc = Document()
    doc.add_paragraph("POSITION SUMMARY")
    doc.add_paragraph("The Strategic Grants Advisor supports grant development.")
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    for duty in _DUTIES:
        add_duty(doc, duty)  # type: ignore[operator]
    doc.add_paragraph("QUALIFICATIONS")
    doc.add_paragraph("Bachelor's degree and five years of related experience.")
    return _blob(doc)


def _duties_of(blob: bytes) -> list[SFUDuty]:
    return parse_jd(extract_text(blob, DocumentFormat.DOCX)).jd.duties


# ── Defect 1: Word list paragraphs ───────────────────────────────────────────


def test_direct_numpr_list_items_parse_as_separate_duties() -> None:
    """🔴 THE REGRESSION. Six ``ListParagraph`` duties with a direct ``numPr`` parsed as
    ONE duty under v8 (every item appended to the first, then cut at 400)."""
    duties = _duties_of(_jdfn_doc(_list_para))

    assert [d.statement for d in duties] == [d.rstrip(".") + "." for d in _DUTIES]


def test_style_inherited_numbering_parse_as_separate_duties() -> None:
    """The older template numbers through the STYLE — nothing on the paragraph."""
    doc = Document()
    style_id = _numbered_style(doc)
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    for duty in _DUTIES:
        p = doc.add_paragraph(duty)
        p._p.get_or_add_pPr().get_or_add_pStyle().val = style_id
    duties = _duties_of(_blob(doc))

    assert len(duties) == len(_DUTIES)


def test_numid_zero_switches_numbering_off() -> None:
    """``numId=0`` on the paragraph is Word's explicit "not a list", even under a
    numbered style — so those paragraphs keep the v8 reading (no marker)."""
    doc = Document()
    style_id = _numbered_style(doc)
    p = doc.add_paragraph("Not actually a bullet.")
    ppr = p._p.get_or_add_pPr()
    ppr.get_or_add_pStyle().val = style_id
    ppr.append(_numpr("0"))

    text = extract_text(_blob(doc), DocumentFormat.DOCX)
    assert LIST_ITEM_MARKER not in text


def test_typed_bullets_are_unchanged() -> None:
    """A typed ``•`` already splits; the marker is never doubled."""

    def typed(doc: DocxDocument, text: str) -> None:
        doc.add_paragraph(f"•\t{text}")

    blob = _jdfn_doc(typed)
    assert len(_duties_of(blob)) == len(_DUTIES)
    assert f"{LIST_ITEM_MARKER}•" not in extract_text(blob, DocumentFormat.DOCX)


def test_plain_paragraphs_extract_byte_identically() -> None:
    """No list formatting, no change: the marker is the ONLY extraction difference."""
    doc = Document()
    for line in ("POSITION SUMMARY", "Plain prose.", "More prose."):
        doc.add_paragraph(line)
    assert extract_text(_blob(doc), DocumentFormat.DOCX) == (
        "POSITION SUMMARY\nPlain prose.\nMore prose."
    )


def test_a_numbered_heading_still_opens_its_section() -> None:
    """Outline-numbered headings are list paragraphs too. Marking one would hide it
    from the heading matcher and silently merge two sections — so headings (JDFN and
    WJQ vocabularies) are never marked."""
    doc = Document()
    _list_para(doc, "DUTIES AND RESPONSIBILITIES", num_id="9")
    for duty in _DUTIES[:3]:
        _list_para(doc, duty)
    _list_para(doc, "QUALIFICATIONS", num_id="9")
    _list_para(doc, "Bachelor's degree in a related field.")
    result = parse_jd(extract_text(_blob(doc), DocumentFormat.DOCX))

    assert len(result.jd.duties) == 3
    assert [q.text for q in result.jd.qualifications] == [
        "Bachelor's degree in a related field."
    ]


def test_a_wjq_docx_keeps_the_unmarked_reading() -> None:
    """The WJQ duty reader enters "bulleted" mode on ONE bullet and then fuses every
    plain paragraph into the duty before it. On a ``.docx`` that mixes list and plain
    paragraphs that measured net-harmful (183 documents lost duties, up to 37 -> 5), so
    a WJQ questionnaire extracts exactly as before v9."""
    doc = Document()
    doc.add_paragraph("WEIGHTED JOB QUESTIONNAIRE (WJQ) CUSTOM")
    doc.add_paragraph("3.  MAJOR FUNCTIONS")
    _list_para(doc, "Answers enquiries about registration. (D)")
    doc.add_paragraph("Prepares weekly enrolment reports for the Registrar.")
    doc.add_paragraph("Reconciles term statistics at semester end.")
    text = extract_text(_blob(doc), DocumentFormat.DOCX)

    assert LIST_ITEM_MARKER not in text
    assert len(parse_jd(text).jd.duties) == 3


def test_bulleted_contacts_stay_in_their_connections_list() -> None:
    """A bullet under "Internal Connections" is a logical line of its own; before v9 it
    matched no label and was appended to ``supervisory`` (213 documents measured)."""
    doc = Document()
    doc.add_paragraph("RELATIONSHIPS")
    doc.add_paragraph("Establishes and maintains relationships.")
    doc.add_paragraph("Internal Connections")
    _list_para(doc, "Director, Research Services - policy updates")
    _list_para(doc, "Associate Dean - postdoctoral opportunities")
    doc.add_paragraph("External Connections")
    _list_para(doc, "Immigration Canada - policy clarification")
    rel = parse_jd(extract_text(_blob(doc), DocumentFormat.DOCX)).jd.relationships

    assert rel is not None
    assert "Director, Research Services - policy updates" in rel.internal
    assert "Associate Dean - postdoctoral opportunities" in rel.internal
    assert "Immigration Canada - policy clarification" in rel.external
    assert rel.supervisory == "Establishes and maintains relationships."


def test_a_fallback_title_never_carries_the_list_glyph() -> None:
    doc = Document()
    _list_para(doc, "We are a leading research university.")
    title = parse_jd(extract_text(_blob(doc), DocumentFormat.DOCX)).jd.title
    assert title == "We are a leading research university."


def test_list_marker_is_invisible_to_tier2_near_dup_shingles() -> None:
    """Tier-2 near-dup re-extracts raw text (``dedup.text_source: raw_clean``); the
    marker must not move a single shingle, or every Word-list document's near-dup
    edges would shift for a formatting reason."""
    plain = "Writes funding proposals.\nReviews proposal budgets with investigators."
    marked = (
        f"{LIST_ITEM_MARKER}Writes funding proposals.\n"
        f"{SUB_ITEM_MARKER}Reviews proposal budgets with investigators."
    )
    size = get_rules().dedup.shingle_size
    marked_hashes = shingle_hashes(marked, shingle_size=size, redact_boilerplate=False)
    plain_hashes = shingle_hashes(plain, shingle_size=size, redact_boilerplate=False)
    assert len(plain_hashes) > 0
    assert marked_hashes == plain_hashes


# ── Q1 (owner ruling 2026-10-20): a "… by:" headline is the duty, items its how/why ──


def test_typed_by_headlines_group_their_items_as_how_why() -> None:
    """Typed bullets throughout: the headline is the duty (kept, not dropped as at v8)
    and the items under it are its how/why, until the next headline."""
    text = (
        "DUTIES AND RESPONSIBILITIES\n"
        "• Manages complex proposal processes by:\n"
        "• Developing and maintaining a proposals calendar.\n"
        "• Tracking relevant grant applications.\n"
        "• Supports Principal Investigators by:\n"
        "• Reviewing budgets for completeness.\n"
    )
    duties = parse_jd(text).jd.duties

    assert [d.statement for d in duties] == [
        "Manages complex proposal processes by",
        "Supports Principal Investigators by",
    ]
    assert duties[0].how_why == [
        "Developing and maintaining a proposals calendar.",
        "Tracking relevant grant applications.",
    ]
    assert duties[1].how_why == ["Reviewing budgets for completeness."]
    assert duties[0].action_verb == "Manages"


def test_word_list_headline_and_nested_items_group() -> None:
    """The modern template: the headline and its sub-bullets are BOTH Word list
    paragraphs; the sub-bullets sit one level in (a deeper numbering indent)."""
    doc = Document()
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    doc.add_paragraph("Identifies funding opportunities by:", style="List Bullet")
    doc.add_paragraph("Writing strategic funding proposals.", style="List Bullet 2")
    doc.add_paragraph("Developing proposal budgets.", style="List Bullet 2")
    doc.add_paragraph("Liaises with external funders.", style="List Bullet")
    duties = _duties_of(_blob(doc))

    assert [d.statement for d in duties] == [
        "Identifies funding opportunities by",
        "Liaises with external funders.",
    ]
    assert duties[0].how_why == [
        "Writing strategic funding proposals.",
        "Developing proposal budgets.",
    ]


def test_the_j4_export_shape_reads_back_as_duty_plus_how_why() -> None:
    """J4 writes a duty as a TYPED bullet line and its how/why as "List Bullet 2"
    sub-bullets — no colon. Nesting alone must group them."""
    doc = Document()
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    for duty, clauses in (
        ("Coordinates the grants calendar", ["so deadlines are met", "with PIs"]),
        ("Maintains records", ["in the research database"]),
    ):
        item = doc.add_paragraph(f"•\t{duty}")
        item.paragraph_format.left_indent = Inches(0.25)
        for clause in clauses:
            doc.add_paragraph(clause, style="List Bullet 2")
    duties = _duties_of(_blob(doc))

    assert [(d.statement, d.how_why) for d in duties] == [
        ("Coordinates the grants calendar", ["so deadlines are met", "with PIs"]),
        ("Maintains records", ["in the research database"]),
    ]


def test_a_list_with_no_headline_stays_separate_duties() -> None:
    doc = Document()
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    for duty in _DUTIES:
        doc.add_paragraph(duty, style="List Bullet")
    duties = _duties_of(_blob(doc))

    assert len(duties) == len(_DUTIES)
    assert all(d.how_why == [] for d in duties)


def test_another_intro_line_ends_the_group_and_is_kept() -> None:
    """Only the template's "… by:" is a headline. Any other "…:" intro closes the group
    and is KEPT as a line of its own (it was dropped before v9 — but v8 had fused it
    into Word-list run-ons, so dropping it now lost words in 303 documents)."""
    text = (
        "DUTIES AND RESPONSIBILITIES\n"
        "• Manages proposals by:\n"
        "• Tracking deadlines.\n"
        "• Budget and Fixed Asset Management:\n"
        "• Orders office supplies.\n"
    )
    duties = parse_jd(text).jd.duties

    assert [(d.statement, d.how_why) for d in duties] == [
        ("Manages proposals by", ["Tracking deadlines."]),
        ("Budget and Fixed Asset Management", []),
        ("Orders office supplies.", []),
    ]


def test_a_headline_fused_onto_a_run_on_keeps_the_run_on() -> None:
    """Unbulleted paragraphs fuse into the item before them, so a "by:" headline can
    arrive as the LAST sentence of a run-on. Only that sentence is the headline."""
    text = (
        "DUTIES AND RESPONSIBILITIES\n"
        "• Researches current web solutions. Conducts pre-release testing of new "
        "features. Maintains the website by:\n"
        "• Editing existing material.\n"
    )
    duties = parse_jd(text).jd.duties

    assert [(d.statement, d.how_why) for d in duties] == [
        (
            "Researches current web solutions. Conducts pre-release testing of new "
            "features.",
            [],
        ),
        ("Maintains the website by", ["Editing existing material."]),
    ]


def test_an_over_cap_run_on_before_a_headline_is_cut_at_sentences_not_lost() -> None:
    """MEASURED (6 documents): a run-on longer than the statement cap, fused before a
    headline, lost its tail to the cap. It is split at sentence breaks instead."""
    sentences = [f"Performs duty number {i} for the unit." for i in range(60)]
    run_on = " ".join(sentences)
    cap = get_rules().segmentation.duty_statement_max_chars
    assert len(run_on) > cap
    text = (
        f"DUTIES AND RESPONSIBILITIES\n• {run_on} Maintains the site by:\n• Editing.\n"
    )
    duties = parse_jd(text).jd.duties

    kept = " ".join(d.statement for d in duties[:-1])
    assert kept == run_on
    assert all(len(d.statement) <= cap for d in duties)
    assert (duties[-1].statement, duties[-1].how_why) == (
        "Maintains the site by",
        ["Editing."],
    )


def test_a_semicolon_run_ending_in_an_intro_keeps_its_duties() -> None:
    """MEASURED on the copy: a long semicolon run whose LAST part is an intro
    ("… including:") was dropped whole by the first grouped build (148 documents lost
    duty words). Each part is placed on its own; the intro part is kept as a line."""
    run = (
        "Identifies sponsorship needs for the campus programs each year; prepares "
        "proposals to vendors and community partners for every event; organizes the "
        "annual signature events and engagement programs, including:"
    )
    assert len(run) > 200
    duties = parse_jd(f"DUTIES AND RESPONSIBILITIES\n• {run}\n").jd.duties

    assert [d.statement for d in duties] == [
        "Identifies sponsorship needs for the campus programs each year",
        "prepares proposals to vendors and community partners for every event",
        "organizes the annual signature events and engagement programs, including",
    ]


def test_a_long_how_why_clause_is_kept_whole() -> None:
    clause = (
        "Reviewing " + "the financial and compliance aspects of applications, " * 12
    )
    text = f"DUTIES AND RESPONSIBILITIES\n• Supports applicants by:\n• {clause}\n"
    duty = parse_jd(text).jd.duties[0]

    assert len(clause.strip()) > 400
    assert duty.how_why == [clause.strip()]


def test_the_how_why_caps_and_headline_are_rulebook_data(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import re

    from src.jd_core import rules as rules_mod

    real = get_rules()
    tuned = real.model_copy(
        update={
            "segmentation": real.segmentation.model_copy(
                update={
                    "duty_how_why_max_count": 1,
                    "duty_how_why_max_chars": 5,
                    "duty_headline_pattern": re.compile(r"(?i)\bthrough\s*:\s*$"),
                }
            )
        }
    )
    monkeypatch.setattr(rules_mod, "get_rules", lambda: tuned)
    text = (
        "DUTIES AND RESPONSIBILITIES\n"
        "• Recruits students through:\n"
        "• Presentations on campus.\n"
        "• Open houses.\n"
    )
    duty = parse_jd(text).jd.duties[0]

    assert duty.statement == "Recruits students through"
    assert duty.how_why == ["Prese"]


# ── Review H1: "…:" lead-in lines outside DUTIES are content, not furniture ──────


def test_a_qualification_ending_in_a_colon_is_kept() -> None:
    """MEASURED (review H1): reading Word lists exposed ":"-ending lines that v8 had
    fused into run-ons, and the qualifications reader dropped them — e.g.
    ``20190808_00121712_Manager.docx`` lost its whole degree requirement."""
    doc = Document()
    doc.add_paragraph("QUALIFICATIONS")
    _list_para(
        doc,
        "Bachelor's degree in Communications, Marketing or a related field plus five "
        "years of related experience OR a combination of equivalent education, "
        "experience, and training:",
    )
    _list_para(doc, "Excellent written communication skills.")
    quals = parse_jd(extract_text(_blob(doc), DocumentFormat.DOCX)).jd.qualifications

    assert [q.text for q in quals] == [
        "Bachelor's degree in Communications, Marketing or a related field plus five "
        "years of related experience OR a combination of equivalent education, "
        "experience, and training",
        "Excellent written communication skills.",
    ]


def test_a_decision_making_lead_in_is_kept() -> None:
    """The template's "The <title> makes decisions related to:" lead-in, and
    lead-ins carrying real scope, were dropped (4,784 documents lost 39,101 Decision
    Making words vs v8; 2,231 / 14,686 in Problem Solving)."""
    text = (
        "IMPACT OF DECISION MAKING\n"
        "• The Strategic Grants Advisor makes decisions related to:\n"
        "• Funding opportunities and related proposals.\n"
        "PROBLEM SOLVING AND LEVEL OF SUPERVISION\n"
        "• Solves complex problems for the Vice-President, Research, including:\n"
        "• Grant deadlines and submissions.\n"
    )
    jd = parse_jd(text).jd

    assert jd.decision_making == [
        "The Strategic Grants Advisor makes decisions related to",
        "Funding opportunities and related proposals.",
    ]
    assert jd.problem_solving == [
        "Solves complex problems for the Vice-President, Research, including",
        "Grant deadlines and submissions.",
    ]


def test_list_glyphs_never_spend_the_summary_cap() -> None:
    """MEASURED: 40 summaries already at the 4,000 cap lost 164 tail words because the
    extractor's list glyphs took up characters. The summary is prose — no glyphs."""
    doc = Document()
    doc.add_paragraph("POSITION SUMMARY")
    for i in range(3):
        _list_para(doc, f"Supports program area number {i} across the campus.")
    summary = parse_jd(
        extract_text(_blob(doc), DocumentFormat.DOCX)
    ).jd.position_summary

    assert summary is not None and "•" not in summary
    assert summary.startswith("Supports program area number 0")


# ── v9-Bank smoke findings (2026-10-20) ──────────────────────────────────────────


def test_a_typed_number_inside_a_nested_list_stays_a_how_why_clause() -> None:
    """FOUND by smoke on the v9 Bank (export -> re-read, 2 roles): a how/why clause
    that itself starts with a typed number ("2. Calculating leases …", "(2) support …")
    was read as a TOP-level item, so every later clause became a duty of its own
    (8 duties exported, 13 re-read). A nested Word list paragraph is a sub-item
    whatever its first characters are."""
    doc = Document()
    doc.add_paragraph("DUTIES AND RESPONSIBILITIES")
    item = doc.add_paragraph("•\tManages campus leases by")
    item.paragraph_format.left_indent = Inches(0.25)
    for clause in (
        "Acting as the point of contact with landlords.",
        "2. Calculating leases and GST rebates.",
        "Collecting monthly rent revenue.",
    ):
        doc.add_paragraph(clause, style="List Bullet 2")
    duties = _duties_of(_blob(doc))

    assert [(d.statement, len(d.how_why)) for d in duties] == [
        ("Manages campus leases by", 3)
    ]
    assert "2. Calculating leases and GST rebates." in duties[0].how_why


def test_a_title_followed_by_a_position_number_label_is_cut_there() -> None:
    """``NEXT_LABEL_RX`` ended in ``\\b``, which cannot match after "#", so
    "Position #:" never cut a label value: 80 archive titles read as
    "Director, Recreational Services & Athletics Position #:" (v8 and v9 alike)."""
    text = (
        "Position Title: Director, Recreational Services & Athletics Position #:\n"
        "POSITION SUMMARY\nLeads the recreation programs.\n"
    )
    assert parse_jd(text).jd.title == "Director, Recreational Services & Athletics"


def test_a_position_number_on_the_title_line_is_read_not_lost() -> None:
    """Once the title is cut at "Position #:", the number after it must land in
    ``position_number`` — measured: 92 archive titles carried one."""
    jd = parse_jd(
        "Position Title: Assistant Football Coach Position #: 204\n"
        "POSITION SUMMARY\nCoaches the team.\n"
    ).jd
    assert (jd.title, jd.position_number) == ("Assistant Football Coach", "204")
    bundled = parse_jd(
        "Position Title: System Administrator Position #: 00129644, 00130935\n"
    ).jd
    assert (bundled.title, bundled.position_number) == (
        "System Administrator",
        "00129644, 00130935",
    )
    rtf = parse_jd(
        "Position Title: |Specialist, Career Education|Position #:|00133267|\n"
    ).jd
    assert rtf.position_number == "00133267"


def test_a_hash_touching_its_number_is_still_cut_from_the_title() -> None:
    """Review round 3: ``(?!\\w)`` after "#" cannot match when the id touches it
    ("Position #110611"), which the old ``\\b`` did cut — 5 archive titles (2
    positions, e.g. 20111017_00110611_Coord_Student_Life.doc) GAINED "Position
    #110611" at 18a2792. The title is cut there and the id is the position number."""
    jd = parse_jd(
        "Position Title:  Coordinator, Student Life         Position #110611\n"
        "Name:\nPOSITION SUMMARY\nSupports students.\n"
    ).jd
    assert (jd.title, jd.position_number) == ("Coordinator, Student Life", "110611")


def test_a_fallback_title_has_its_whitespace_collapsed() -> None:
    """A fallback title is a raw line, and antiword pads columns with runs of spaces
    ("HUMAN RESOURCES             #01449"); every other title is cleaned."""
    assert parse_jd("HUMAN RESOURCES             #01449\n").jd.title == (
        "HUMAN RESOURCES #01449"
    )


def test_a_stretched_wjq_marker_still_routes_to_the_wjq_reader() -> None:
    """antiword stretches a heading's words apart ("Weighted   Job   Questionnaire");
    the primary marker was matched as a plain substring, so the CUPE form went down
    the JDFN path and its banner became the title."""
    text = (
        "Weighted   Job   Questionnaire   (WJQ)   Custom\n"
        "1.  POSITION IDENTIFICATION\nDepartment Position Title:\nClerk\n"
    )
    assert parse_jd(text).template == "wjq"


def test_a_typed_bullet_with_no_space_still_splits() -> None:
    text = "DUTIES AND RESPONSIBILITIES\n•Tracks budgets.\n•Reviews invoices.\n"
    assert [d.statement for d in parse_jd(text).jd.duties] == [
        "Tracks budgets.",
        "Reviews invoices.",
    ]


def test_the_bullet_count_cap_is_rulebook_data() -> None:
    """``_MAX_BULLETS`` (20) cut Decision Making / Problem Solving lists; it is now
    ``segmentation.bullet_max_count`` (HR-268), under the model's own ceiling."""
    seg = get_rules().segmentation
    assert (
        20 < seg.bullet_max_count <= _max_length(SFUJobDescription, "decision_making")
    )
    many = "\n".join(f"• Decides on matter number {i}." for i in range(30))
    jd = parse_jd(f"IMPACT OF DECISION MAKING\n{many}\n").jd
    assert len(jd.decision_making) == 30


# ── Defect 2: the 400-character duty statement ───────────────────────────────


def test_a_long_duty_statement_survives_past_400() -> None:
    long_duty = "Coordinates " + "the research grant lifecycle and reporting, " * 20
    assert len(long_duty) > 800
    result = parse_jd(f"DUTIES AND RESPONSIBILITIES\n• {long_duty.strip()}\n")

    assert len(result.jd.duties[0].statement) == len(long_duty.strip())


def test_the_statement_cap_is_rulebook_data_and_fits_the_contract() -> None:
    seg = get_rules().segmentation
    ceiling = _max_length(SFUDuty, "statement")
    assert 400 < seg.duty_statement_max_chars <= ceiling


def test_the_caps_are_applied_from_the_rulebook(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Retune the knobs and the parser follows — no code constant in the way."""
    from src.jd_core import rules as rules_mod

    real = get_rules()
    tuned = real.model_copy(
        update={
            "segmentation": real.segmentation.model_copy(
                update={"duty_statement_max_chars": 50, "duty_max_count": 2}
            )
        }
    )
    monkeypatch.setattr(rules_mod, "get_rules", lambda: tuned)
    long_duty = "Coordinates " + "x " * 100
    result = parse_jd(
        f"DUTIES AND RESPONSIBILITIES\n• {long_duty}\n• Second duty.\n• Third duty.\n"
    )
    assert len(result.jd.duties) == 2
    assert len(result.jd.duties[0].statement) == 50


# ── Defect 3: the 12-function WJQ ceiling ────────────────────────────────────

_WJQ_HEAD = """SIMON FRASER UNIVERSITY & C.U.P.E., LOCAL 3338
WEIGHTED JOB QUESTIONNAIRE (WJQ) CUSTOM
1.  POSITION IDENTIFICATION
Department Position Title:
Program Assistant
3.  MAJOR FUNCTIONS
"""


def _wjq_with(major: int, minor: int) -> str:
    lines = [_WJQ_HEAD.rstrip("\n")]
    lines += [f"• Performs major function number {i}. (D)" for i in range(1, major + 1)]
    lines.append("4.  MINOR FUNCTIONS")
    lines += [f"• Performs minor function number {i}." for i in range(1, minor + 1)]
    lines.append("13.  QUALIFICATIONS")
    lines.append("High School graduation.")
    return "\n".join(lines) + "\n"


def test_a_fourteen_function_cupe_questionnaire_keeps_all_fourteen() -> None:
    """HR's own Dec-2024 CUPE sample carries 14 functions; v8 kept 12 (majors first),
    so the last minor functions vanished."""
    result = segment_wjq(_wjq_with(major=10, minor=4), get_rules().wjq)

    assert len(result.jd.duties) == 14
    assert result.jd.duties[-1].statement == "Performs minor function number 4"


def test_a_full_24_function_questionnaire_fits() -> None:
    """The form carries 12 major + 12 minor slots, and the Builder exports that many
    (HR-246) — a re-read must be able to hold what the Builder can write."""
    result = segment_wjq(_wjq_with(major=12, minor=12), get_rules().wjq)
    assert len(result.jd.duties) == 24


def test_the_duty_count_caps_are_rulebook_data_and_fit_the_contract() -> None:
    seg = get_rules().segmentation
    ceiling = _max_length(SFUJobDescription, "duties")
    assert 12 < seg.duty_max_count <= ceiling
    assert 24 <= seg.wjq_duty_max_count <= ceiling


def test_the_parser_version_moved() -> None:
    assert PARSER_VERSION == "jd_segmenter_v9"


def _max_length(model: type[object], field: str) -> int:
    fields = model.model_fields  # type: ignore[attr-defined]
    return int(
        next(
            m.max_length
            for m in fields[field].metadata
            if getattr(m, "max_length", None) is not None
        )
    )


@pytest.mark.parametrize(
    ("knob", "model", "field"),
    [
        ("duty_statement_max_chars", SFUDuty, "statement"),
        ("duty_how_why_max_chars", SFUDuty, "statement"),
        ("duty_how_why_max_count", SFUDuty, "how_why"),
        ("duty_max_count", SFUJobDescription, "duties"),
        ("wjq_duty_max_count", SFUJobDescription, "duties"),
    ],
)
def test_a_cap_above_the_model_ceiling_is_refused_at_load(
    knob: str, model: type[object], field: str
) -> None:
    """Review (v9): the loader's ceiling check had no test of its own — removing it
    left the suite green, because the shipped values comply. Over the ceiling a long
    document is a parse FAILURE, not a trim, so the rulebook must refuse to load."""
    from pathlib import Path

    import yaml
    from pydantic import ValidationError

    import src.jd_core.rules as rules_pkg
    from src.jd_core.rules import Segmentation

    raw = Path(rules_pkg.__file__).with_name("segmentation.yaml").read_text("utf-8")
    data = yaml.safe_load(raw)
    Segmentation.model_validate(data)  # the shipped file loads
    data[knob] = _max_length(model, field) + 1
    with pytest.raises(ValidationError, match=knob):
        Segmentation.model_validate(data)
