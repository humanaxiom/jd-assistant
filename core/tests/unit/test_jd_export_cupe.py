"""Track J, PR J1 — a CUPE draft exports on HR's CUPE form, not the APSA layout.

CUPE 3338 Art. 15.04: the standard job description form "cannot be revised without
mutual agreement between the University and the Union". Before J1 every draft, CUPE
included, was rendered in the APSA/JDFN layout with the APSA territorial/equity footer.

The central property is the ROUND TRIP: a CUPE JD exported here and read back by the
Bank's own ingest path (``extract_text`` -> ``parse_jd``) comes back with everything
the WJQ parser can read (oracle: ``tests/export_roundtrip.py``).
"""

from __future__ import annotations

import zipfile
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.api.main import app
from src.jd_bank.composer.wjq_answers import WJQAnswers
from src.jd_bank.composer.wjq_assemble import assemble_wjq_jd
from src.jd_bank.db.models import DocumentFormat
from src.jd_bank.ingest.extract import extract_text
from src.jd_core.models.parsed_jd import (
    JobClassification,
    SFUDuty,
    SFUJobDescription,
    SFUQualification,
    SFURelationships,
)
from src.jd_core.parser.segmenter import parse_jd
from src.jd_core.quality.validators import template_of
from src.jd_core.rules import Rules, get_rules
from src.jd_export import render_cupe_docx, render_jdfn_docx, render_sfu_docx
from tests.export_roundtrip import cupe_round_trip_losses


@pytest.fixture(scope="module")
def rules() -> Rules:
    return get_rules()


def _context(rules: Rules, bodies: dict[str, str]) -> str:
    """An ``additional_context`` written the way the parser and assembler write it:
    each point-factor block under its canonical heading."""
    headings = rules.wjq.section_headings
    return "\n\n".join(f"{headings[k][0]}\n{v}" for k, v in bodies.items())


def _cupe_jd(rules: Rules, **update: object) -> SFUJobDescription:
    jd = SFUJobDescription(
        title="Library Circulation Assistant",
        department="Belzberg Library",
        position_number="12345, 12346",
        grade="6",
        classification=JobClassification(scheme="cupe", value="6", source="parsed"),
        employee_group="cupe",
        position_summary=(
            "Provides front-line circulation services so patrons can reach the "
            "collection. Processes loans, returns and patron accounts."
        ),
        duties=[
            SFUDuty(
                action_verb="Circulate",
                statement="library materials by recording transactions and shelving",
                how_why=["to ensure timely patron access", "using the ILS"],
                frequency="daily",
            ),
            SFUDuty(
                action_verb="Processes",
                statement="Processes interlibrary deliveries (20%)",
                frequency="weekly",
            ),
            SFUDuty(
                action_verb="Compiles",
                statement="Compiles annual activity statistics",
                frequency=None,
            ),
            SFUDuty(
                action_verb="Reviews",
                statement="Reviews reserve lists each term",
                frequency="semester",
            ),
        ],
        relationships=SFURelationships(
            supervisory="Temporary circulation staff",
            internal=["Library staff", "Faculty"],
            external=["Vendors", "General public"],
        ),
        qualifications=[
            SFUQualification(text="High school graduation", kind="education"),
            SFUQualification(
                text="2 years related library experience", kind="experience"
            ),
            SFUQualification(
                text="Knowledge of Library of Congress classification", kind="knowledge"
            ),
            SFUQualification(
                text="keyboarding at 45 wpm", kind="skill", modifier="intermediate"
            ),
            SFUQualification(text="Ability to push loaded book trucks", kind="ability"),
            SFUQualification(text="Criminal record check", kind="security"),
        ],
        additional_context=_context(
            rules,
            {
                "level_of_independence": "Works from general instructions.",
                "training_exercised": "Explains own work to replacement staff.",
                "direction_exercised": "Checks the work of student assistants.",
                "impact_of_errors": "Misfiled items delay patron access.",
                "effort": "Visual: reading call numbers, 1-2 hours, frequent.",
                "working_conditions": "General office 35 / 35 hours.",
                "continuing_education": "Learns new circulation procedures.",
            },
        ),
    )
    return jd.model_copy(update=update)


def _texts(data: bytes) -> list[str]:
    """Every paragraph's text in document order, table cells included."""
    doc = Document(BytesIO(data))
    out: list[str] = []

    def walk(container) -> None:  # type: ignore[no-untyped-def]
        for item in container.iter_inner_content():
            if hasattr(item, "rows"):
                for row in item.rows:
                    seen: list[object] = []
                    for cell in row.cells:
                        if cell._tc in seen:  # merged cells repeat
                            continue
                        seen.append(cell._tc)
                        walk(cell)
            else:
                out.append(item.text)

    walk(doc)
    return out


def _members(data: bytes) -> list[tuple[str, bytes]]:
    """A ``.docx`` as its zip members' names and CONTENTS. Raw bytes are not
    comparable: python-docx stamps each zip entry with the wall-clock second, so two
    renders a second apart differ in bytes while being the same document."""
    archive = zipfile.ZipFile(BytesIO(data))
    return [(i.filename, archive.read(i.filename)) for i in archive.infolist()]


def _reparse(data: bytes) -> SFUJobDescription:
    result = parse_jd(extract_text(data, DocumentFormat.DOCX))
    assert result.template == "wjq", "the export is not recognised as a WJQ form"
    return result.jd


# ── dispatch: one entry point, the form follows the JD ──────────────────────────


def test_a_cupe_jd_exports_on_the_cupe_form(rules: Rules) -> None:
    jd = _cupe_jd(rules)
    assert template_of(jd) == "wjq"
    texts = _texts(render_sfu_docx(jd))
    for line in rules.cupe_form.banner:
        assert line in texts, f"CUPE banner line {line!r} missing"
    # ...and not the APSA layout's headings.
    assert "Duties and Responsibilities" not in texts
    assert "Relationships" not in texts


def test_the_single_entry_point_routes_cupe_to_the_cupe_renderer(rules: Rules) -> None:
    jd = _cupe_jd(rules)
    assert _members(render_sfu_docx(jd)) == _members(render_cupe_docx(jd))


def test_a_jdfn_jd_is_rendered_exactly_as_before(rules: Rules) -> None:
    """APSA output is untouched: the dispatcher hands a non-CUPE JD to the JDFN
    renderer, byte for byte."""
    jd = _cupe_jd(rules, employee_group="apsa", additional_context=None)
    assert _members(render_sfu_docx(jd)) == _members(render_jdfn_docx(jd))
    texts = _texts(render_sfu_docx(jd))
    assert "Duties and Responsibilities" in texts


def test_the_cupe_export_carries_no_apsa_footer(rules: Rules) -> None:
    joined = "\n".join(_texts(render_sfu_docx(_cupe_jd(rules))))
    boiler = rules.boilerplate
    for passage in (
        *boiler.territorial_acknowledgement,
        *boiler.employment_equity,
        *boiler.about_sfu,
    ):
        assert passage not in joined, "APSA boilerplate leaked onto the CUPE form"
    assert "Territorial Acknowledgement" not in joined


def test_the_fourteen_sections_render_in_the_forms_order(rules: Rules) -> None:
    texts = _texts(render_sfu_docx(_cupe_jd(rules)))
    positions = []
    for number, section in enumerate(rules.wjq.section_headings, start=1):
        heading = rules.wjq.section_headings[section][0]
        hits = [
            i
            for i, t in enumerate(texts)
            if t.strip().startswith(f"{number}.") and heading in t
        ]
        assert hits, f"section {number} {heading} missing"
        positions.append(hits[0])
    assert positions == sorted(positions), "sections out of HR's order"


# ── NN #1: nothing on the form asserts an approval ─────────────────────────────


def test_the_hr_box_and_the_sign_off_lines_are_blank(rules: Rules) -> None:
    """The draft's grade is never printed — not in HR's box, not anywhere — and the
    Approval and Review table holds only HR's own labels."""
    data = render_sfu_docx(
        _cupe_jd(rules, grade="Grade 7 secret", position_number=None)
    )
    texts = _texts(data)
    assert not any("Grade 7 secret" in t for t in texts)

    doc = Document(BytesIO(data))
    form = rules.cupe_form
    approval = doc.tables[-1]
    allowed = {
        "",
        f"14.  {rules.wjq.section_headings['approval_review'][0]}",
        form.approval_supervisor_statement,
        form.approval_employee_statement,
        *form.approval_supervisor_labels,
        *form.approval_employee_labels,
    }
    for row in approval.rows:
        for cell in row.cells:
            assert (
                cell.text.strip() in allowed
            ), f"sign-off cell pre-filled: {cell.text!r}"

    id_table = doc.tables[0]
    hr_values = []
    for row in id_table.rows:
        cells = row.cells
        for i, cell in enumerate(cells):
            if cell.text.strip() in form.hr_box_labels and i + 1 < len(cells):
                hr_values.append(cells[i + 1].text.strip())
    assert hr_values and all(v == "" for v in hr_values), hr_values


def test_every_checkbox_option_is_printed_unticked(rules: Rules) -> None:
    texts = _texts(render_sfu_docx(_cupe_jd(rules)))
    form = rules.cupe_form
    for section, options in form.checkbox_options.items():
        for option in options:
            # The box, empty, then HR's option text — never a ticked glyph.
            assert f"{form.checkbox_glyph} {option}" in texts, f"{section}: {option!r}"
    assert not any("☒" in t or "☑" in t for t in texts), "a box is ticked"


def test_the_forms_instruction_is_not_printed_twice_but_a_typed_tick_survives(
    rules: Rules,
) -> None:
    """An archive-parsed CUPE role stores its point-factor sections WITH the form's own
    instruction and options flattened into them, and an archive form records its choice
    as a tick TYPED beside an option ("_X …"). The instruction is the form's to print,
    so it is not printed again; the option stays in the draft's text, because removing
    it would orphan the tick — the only record of the choice."""
    form = rules.cupe_form
    options = form.checkbox_options["level_of_independence"]
    stored = " ".join(
        [
            *form.instructions["level_of_independence"],
            options[0],
            f"_X {options[1]}",
            *options[2:],
        ]
    )
    jd = _cupe_jd(
        rules,
        additional_context=_context(
            rules, {"level_of_independence": f"{stored} Checks student work."}
        ),
    )
    texts = _texts(render_cupe_docx(jd))
    instruction = form.instructions["level_of_independence"][0]
    assert sum(instruction in t for t in texts) == 1
    assert any(f"_X {options[1]}" in t for t in texts), "the typed tick was orphaned"
    assert any("Checks student work." in t for t in texts)


def test_a_short_form_heading_is_never_cut_out_of_a_sentence(rules: Rules) -> None:
    """MEASURED on the live Bank: an older form's prompt reads "Describe any
    uncomfortable conditions in the work place(s)". Stripping the column heading
    "Uncomfortable Conditions" anywhere it occurred turned that into "Describe any in
    the work place(s)"."""
    older = "(b). Describe any uncomfortable conditions in the work place(s). Dust."
    jd = _cupe_jd(
        rules, additional_context=_context(rules, {"working_conditions": older})
    )
    joined = "\n".join(_texts(render_cupe_docx(jd)))
    assert "Describe any uncomfortable conditions in the work place(s)" in joined


# ── the round trip through the Bank's own parser ──────────────────────────────


def test_a_cupe_export_round_trips_through_the_wjq_parser(rules: Rules) -> None:
    jd = _cupe_jd(rules)
    back = _reparse(render_sfu_docx(jd))
    assert cupe_round_trip_losses(jd, back, rules) == []
    # The grade is deliberately NOT exported (HR-240), so it cannot come back.
    assert back.grade is None


def test_an_authored_wjq_draft_round_trips(rules: Rules) -> None:
    """The Builder's own CUPE draft — assembled from answers — survives export and
    re-ingest.

    Since J2 this is a PRE-J2 answer set (the shape J1 authored: %-allocation,
    internal / external lines, knowledge / skills / abilities), loaded the way an open
    browser tab's ``answers_json`` would be — so it also pins that such a draft still
    loads, assembles and round-trips (the J2 structure has its own suite,
    ``test_cupe_form_structure.py``)."""
    answers = WJQAnswers.model_validate(
        {
            "title": "Department Secretary",
            "department": "Faculty of Science",
            "position_number": "40001",
            "position_summary": "Provides secretarial support to the department.",
            "major_functions": [
                {
                    "action_verb": "Types",
                    "statement": "Types correspondence",
                    "allocation": 60,
                },
                {"action_verb": "Maintains", "statement": "Maintains office files"},
            ],
            "minor_functions": [
                {"action_verb": "Purges", "statement": "Purges and archives records"}
            ],
            "internal": ["Faculty", "Students"],
            "external": ["Vendors"],
            "education": ["High school graduation"],
            "experience": ["2 years office experience"],
            "knowledge": [{"text": "office procedures", "modifier": "working"}],
            "skills": [{"text": "keyboarding at 55 wpm", "modifier": "intermediate"}],
            "abilities": ["Ability to work occasional evenings"],
            "level_of_independence": "Works from general instructions.",
            "impact_of_errors": "Typing errors are caught by the reader.",
            "working_conditions": "General office, 35 hours.",
        }
    )
    jd = assemble_wjq_jd(answers)
    back = _reparse(render_sfu_docx(jd))
    assert cupe_round_trip_losses(jd, back, rules) == []


def test_the_round_trip_oracle_catches_a_lost_duty(rules: Rules) -> None:
    """The oracle must be able to fail: drop a duty from what comes back."""
    jd = _cupe_jd(rules)
    back = _reparse(render_sfu_docx(jd))
    damaged = back.model_copy(update={"duties": back.duties[:-1]})
    assert cupe_round_trip_losses(jd, damaged, rules)


def test_a_bare_cupe_jd_renders(rules: Rules) -> None:
    jd = SFUJobDescription(title="Bare Role", employee_group="cupe")
    back = _reparse(render_sfu_docx(jd))
    assert back.title == "Bare Role"


# ── the form's labels are the parser's labels ─────────────────────────────────


def test_the_modelled_id_labels_are_ones_the_parser_reads(rules: Rules) -> None:
    form = rules.cupe_form
    for field in ("title", "department", "position_number"):
        label = form.id_labels[field].rstrip(":").strip().lower()  # type: ignore[index]
        known = {x.rstrip(":").strip().lower() for x in rules.wjq.id_labels[field]}  # type: ignore[index]
        assert label in known, f"{field}: the parser cannot read {label!r}"


def test_the_export_route_serves_a_cupe_draft_on_the_cupe_form(rules: Rules) -> None:
    resp = TestClient(app).post(
        "/jd-bank/compose/export",
        json={"title": "Library Assistant", "employee_group": "cupe"},
    )
    assert resp.status_code == 200
    texts = _texts(resp.content)
    assert rules.cupe_form.banner[1] in texts
