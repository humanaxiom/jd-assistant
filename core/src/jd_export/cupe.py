"""Render a CUPE (WJQ) JD on HR's CUPE form — Track J, PR J1.

HR's "WJQ Custom — Part 1: Job Description" is a different form from the APSA/APEX/Poly
template, and CUPE 3338 Art. 15.04 says it "cannot be revised without mutual agreement
between the University and the Union". So a CUPE draft does not export in the APSA
layout, and never carries the APSA territorial/equity footer: this module prints HR's
form — banner, fourteen numbered sections in HR's order, its labels, instructions,
checkbox options and tables — and fills in only what the Bank's model actually holds.

**Everything the form prints that is not the draft's own content is rulebook data**
(``cupe_form.yaml``, CLAUDE.md non-negotiable #2), and the fourteen section HEADINGS are
``wjq.section_headings`` — the words the Bank's own WJQ parser matches — so an export
read back through ``extract_text`` -> ``parse_jd`` lands every field where it came from.
That round trip is the contract (``tests/export_roundtrip.py``; live in the smoke).

**Track J2 fills the form's structure** for a draft AUTHORED on the Builder
(``jd.wjq_form``): the author's identification fields and "completed by" tick, the
minor functions under MINOR FUNCTIONS, the ticked levels of sections 5/6/7/12
(``checkbox_ticked_glyph``, HR-250), the duration / frequency / hours of the 8/10/11
tables, and section 13's certification, number of years, equivalency tick + reason and
Occupational Requirement(s).

What is deliberately NOT filled in (registered in HR-239 / HR-240):

* the **Approval and Review** lines and HR's **"FOR USE BY HUMAN RESOURCES"** box are
  always blank, and the draft's grade is printed nowhere (non-negotiable #1) — the
  author's CURRENT classification prints in its own cell, never in HR's box;
* for a role PARSED from the archive (no ``wjq_form``) the checkbox sections print
  every option unticked and the tables print duration, frequency and hours blank — a
  tick there would be a guess — and every duty prints under MAJOR FUNCTIONS (the parser
  merges sections 3 and 4). The free text is printed inside its section, minus the
  form's own instruction / prompt sentences (an archive-parsed role stores them
  flattened into the section) — but WITH its options and column headings, because an
  archive form types its tick beside an option and that tick is the only record of the
  choice.

Deterministic in CONTENT and pure of I/O beyond building bytes (the raw bytes are not
reproducible across seconds: python-docx stamps each zip entry with the wall clock —
true of the APSA export too). Nothing here validates or publishes.
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from io import BytesIO
from typing import cast, get_args

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from src.jd_core.models.parsed_jd import (
    SFUDuty,
    SFUJobDescription,
    SFUQualification,
    WJQForm,
    WJQRow,
)
from src.jd_core.rules import CupeForm, Rules, Wjq, WjqSection, get_rules

__all__ = ["render_cupe_docx"]

_TABLE_STYLE = "Table Grid"


# ── text helpers ──────────────────────────────────────────────────────────────


def _phrase_rx(phrase: str) -> re.Pattern[str]:
    """``phrase`` matched across ANY whitespace run (the stored text of a parsed form
    has its line breaks collapsed, and a header cell may carry a line break)."""
    return re.compile(r"\s+".join(re.escape(w) for w in phrase.split()), re.IGNORECASE)


def _without(text: str, phrases: Iterable[str]) -> str:
    """``text`` with every occurrence of the form's own INSTRUCTION / PROMPT sentences
    removed — the form prints them itself — and the blank lines that leaves dropped.

    🔴 Only whole instruction and prompt sentences, never checkbox options or column
    headings. MEASURED over the live Bank: stripping the column heading "Uncomfortable
    Conditions" cut it out of an older form's own sentence ("Describe any uncomfortable
    conditions in the work place(s)" became "Describe any in the work place(s)"), and an
    archive form records its tick as typed text BESIDE the option ("_X An equivalent
    combination …") — strip the option and the tick is orphaned, which loses the one
    choice the document recorded. A repeated option reads as redundancy; a lost tick is
    a lost fact. Longest phrase first, so a phrase containing another goes whole.

    Each phrase is removed ONCE — the once the form prints it. A stored section that
    repeats the instruction (a page break in the source form) keeps the repeat."""
    for phrase in sorted({p for p in phrases if p.strip()}, key=len, reverse=True):
        text = _phrase_rx(phrase).sub(" ", text, count=1)
    lines = (re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)


def _split_at(text: str, anchors: Sequence[str]) -> list[str]:
    """``text`` cut at each anchor phrase, in order: ``len(anchors)`` chunks, chunk *i*
    holding what follows anchor *i*. An anchor that is not found contributes an empty
    chunk and its text stays with the chunk before it, and anything before the first
    anchor joins the first chunk — so nothing is ever dropped."""
    chunks = [""] * len(anchors)
    cuts: list[tuple[int, int, int]] = []  # (anchor index, start, end)
    pos = 0
    for i, anchor in enumerate(anchors):
        match = _phrase_rx(anchor).search(text, pos)
        if match is not None:
            cuts.append((i, match.start(), match.end()))
            pos = match.end()
    if not cuts:
        chunks[0] = text
        return chunks
    lead = text[: cuts[0][1]]
    for n, (i, _, end) in enumerate(cuts):
        stop = cuts[n + 1][1] if n + 1 < len(cuts) else len(text)
        chunks[i] = text[end:stop]
    chunks[cuts[0][0]] = f"{lead}\n{chunks[cuts[0][0]]}"
    return chunks


def _context_sections(context: str | None, wjq: Wjq) -> dict[WjqSection, str]:
    """``additional_context`` back into its point-factor sections, by the canonical
    heading line the parser and the assembler both write. Text before the first heading
    belongs to the first point-factor section (the clone transform's rule) rather than
    to nowhere."""
    canonical = {wjq.section_headings[s][0]: s for s in wjq.context_sections}
    current: WjqSection = wjq.context_sections[0]
    blocks: dict[WjqSection, list[str]] = {}
    for line in (context or "").splitlines():
        if line.strip() in canonical:
            current = canonical[line.strip()]
            blocks.setdefault(current, [])
            continue
        blocks.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in blocks.items() if "\n".join(v).strip()}


# ── docx helpers ──────────────────────────────────────────────────────────────


def _bold(paragraph: Paragraph, text: str) -> None:
    paragraph.add_run(text).bold = True


def _set(cell: _Cell, text: str, *, bold: bool = False) -> None:
    """Put ``text`` in ``cell`` — one paragraph per line, so a multi-line value stays a
    sequence of lines rather than one run with breaks."""
    lines = text.split("\n") if text else [""]
    first = cell.paragraphs[0]
    for i, line in enumerate(lines):
        paragraph = first if i == 0 else cell.add_paragraph()
        if bold:
            _bold(paragraph, line)
        elif line:
            paragraph.add_run(line)


def _table(doc: DocxDocument, cols: int) -> Table:
    table = cast(Table, doc.add_table(rows=0, cols=cols))
    table.style = _TABLE_STYLE
    return table


def _new_row(table: Table) -> list[_Cell]:
    """A fresh row's cells (python-docx leaves ``add_row`` unannotated)."""
    return list(table.add_row().cells)  # type: ignore[no-untyped-call]


def _merge(first: _Cell, last: _Cell) -> _Cell:
    return cast(_Cell, first.merge(last))


def _full_row(table: Table, text: str, *, bold: bool = False) -> _Cell:
    """A row whose single cell spans the table."""
    cells = _new_row(table)
    cell = _merge(cells[0], cells[-1]) if len(cells) > 1 else cells[0]
    _set(cell, text, bold=bold)
    return cell


def _heading_row(table: Table, heading: str, instructions: Sequence[str]) -> None:
    cell = _full_row(table, heading, bold=True)
    for line in instructions:
        cell.add_paragraph(line)


def _row(table: Table, values: Sequence[str], *, bold: bool = False) -> None:
    for cell, value in zip(_new_row(table), values, strict=False):
        _set(cell, value, bold=bold)


def _text_rows(table: Table, text: str) -> None:
    """The draft's free text for a section, as full-width rows inside its table."""
    for line in text.splitlines():
        if line.strip():
            _full_row(table, line.strip())


# ── field rendering ────────────────────────────────────────────────────────────


def _duty_text(duty: SFUDuty, letters: dict[str, str]) -> str:
    """One duty as HR's form wants it: verb + what, then its how/why clauses, then the
    ``(D)/(W)/(M)/(S)`` frequency code (the letters are ``wjq.frequency_markers``, the
    parser's own, so the code reads back). The verb is prepended only when the
    statement does not already lead with it, so a duty never doubles or drops it."""
    statement = duty.statement.strip()
    verb = duty.action_verb.strip()
    if verb and not statement.lower().startswith(verb.lower()):
        statement = f"{verb} {statement}"
    parts = [statement, *(h.strip() for h in duty.how_why)]
    text = ", ".join(p.rstrip(" .;,") for p in parts if p.strip()) + "."
    if duty.frequency:
        code = letters.get(duty.frequency, duty.frequency)
        text = f"{text} ({code})"
    return text


def _qual_text(qual: SFUQualification) -> str:
    """The qualification, with its Toolkit modifier prefixed when not already in the
    text — the APSA export's rule, so a modifier is never dropped."""
    text = qual.text.strip()
    modifier = (qual.modifier or "").strip()
    if modifier and modifier.lower() not in text.lower():
        return f"{modifier} {text}"
    return text


def _headings(wjq: Wjq) -> dict[WjqSection, str]:
    """``N.  HEADING`` for the fourteen sections, numbered in the form's order, the
    words taken from the parser's own vocabulary (the canonical variant)."""
    return {
        section: f"{number}.  {wjq.section_headings[section][0]}"
        for number, section in enumerate(get_args(WjqSection), start=1)
    }


def _identification(
    doc: DocxDocument, jd: SFUJobDescription, form: CupeForm, heading: str
) -> None:
    """Section 1: the draft's identification on the left, HR's box on the right —
    whose values are ALWAYS blank (HR-240, non-negotiable #1). The author's current
    classification, date and "completed by" come from ``wjq_form`` (J2) and are blank
    for a role that has none; the evaluating supervisor's title is the top-level
    ``jd.supervisor_title`` (J4), the field the review editor writes (review M2)."""
    labels = form.id_labels
    hr = list(form.hr_box_labels)
    extra = jd.wjq_form or WJQForm()
    table = _table(doc, 4)
    top = _new_row(table)
    _set(_merge(top[0], top[1]), heading, bold=True)
    _set(_merge(top[2], top[3]), form.hr_box_heading, bold=True)
    left: list[tuple[str, str]] = [
        (labels["department"], (jd.department or "").strip()),
        (labels["position_number"], (jd.position_number or "").strip()),
        (labels["current_classification"], extra.current_classification or ""),
        (labels["title"], jd.title.strip()),
        # Track J4's top-level field — the one the review editor writes (review M2).
        (labels["supervisor_title"], (jd.supervisor_title or "").strip()),
        (
            labels["date_completed"],
            extra.date_completed.isoformat() if extra.date_completed else "",
        ),
    ]
    # HR's box sits beside rows 1, 4, 5 and 6 on the form.
    right: list[str] = [hr[0], "", "", *hr[1:]]
    for i, (label, value) in enumerate(left):
        cells = _new_row(table)
        _set(cells[0], label, bold=True)
        _set(cells[1], value)
        if i < len(right) and right[i]:
            _set(cells[2], right[i], bold=True)
    # HR's two options print in the form's order (Employee, then Evaluating
    # Supervisor), which is the order of `WJQCompletedBy`.
    chosen = {"employee": 0, "supervisor": 1}.get(extra.completed_by or "", -1)
    for i, option in enumerate(form.completed_by_options):
        cells = _new_row(table)
        if i == 0:
            _set(cells[0], form.completed_by_label, bold=True)
        _set(cells[1], f"{_box(form, i == chosen)} {option}")


def _box(form: CupeForm, ticked: bool) -> str:
    """The glyph for one option: HR-250's tick, or the form's empty box."""
    return form.checkbox_ticked_glyph if ticked else form.checkbox_glyph


def _checkbox_section(
    doc: DocxDocument,
    heading: str,
    instructions: Sequence[str],
    options: Sequence[str],
    form: CupeForm,
    chosen: Sequence[str],
    text: str,
) -> None:
    """Sections 5, 6, 7, 12: every option with its box — TICKED for the levels the
    author chose (``chosen`` letters, A = first option; J2), empty otherwise — then the
    draft's own text for the section with the form's instruction taken out of it
    (options are kept in it — an archive form's tick is typed beside them; see
    :func:`_without`)."""
    table = _table(doc, 1)
    _heading_row(table, heading, instructions)
    for i, option in enumerate(options):
        _full_row(table, f"{_box(form, chr(ord('A') + i) in chosen)} {option}")
    _text_rows(table, _without(text, instructions))
    doc.add_paragraph()


def _contacts(
    doc: DocxDocument, jd: SFUJobDescription, form: CupeForm, heading: str
) -> None:
    """Section 8: one row per contact. WHICH contacts is ``relationships`` (what a
    reviewer edits); each one's duration and frequency come from ``wjq_form.contacts``
    matched by text (J2) — so a contact edited after authoring prints with blank cells
    rather than someone else's duration. Blank for a role with no ``wjq_form``."""
    rel = jd.relationships
    contacts = [
        c.strip() for c in ([*rel.internal, *rel.external] if rel else []) if c.strip()
    ]
    pool = list(jd.wjq_form.contacts) if jd.wjq_form is not None else []
    rows: list[list[str]] = []
    for contact in contacts:
        match = next((r for r in pool if _same(r.item, contact)), None)
        if match is not None:
            pool.remove(match)
        rows.append(
            [
                contact,
                (match.duration or "") if match else "",
                (match.frequency or "") if match else "",
            ]
        )
    table = _table(doc, len(form.contacts_columns))
    _heading_row(
        table, heading, form.instructions.get("internal_external_contacts", ())
    )
    _row(table, form.contacts_columns, bold=True)
    _table_rows(table, rows, form.blank_rows["contacts"])
    doc.add_paragraph()


def _same(a: str, b: str) -> bool:
    return " ".join(a.lower().split()) == " ".join(b.lower().split())


def _table_rows(table: Table, rows: Sequence[Sequence[str]], minimum: int) -> None:
    """Filled rows, padded with blank rows to ``minimum`` — the empty lines HR's form
    offers. Never truncates."""
    width = len(table.columns)
    for values in rows:
        _row(table, [*values, *([""] * (width - len(values)))][:width])
    for _ in range(max(0, minimum - len(rows))):
        _row(table, [""] * width)


def _impact_of_errors(
    doc: DocxDocument, form: CupeForm, heading: str, text: str
) -> None:
    """Section 9: HR's two prompts, each followed by what the draft holds under it. A
    stored section that carries both prompts is split at them; text with neither prompt
    (an authored draft) cannot be split honestly and goes under the first."""
    prompts = [
        form.impact_of_errors_common_prompt,
        form.impact_of_errors_uncommon_prompt,
    ]
    chunks = _split_at(text, prompts)
    table = _table(doc, 1)
    _heading_row(table, heading, ())
    for prompt, chunk in zip(prompts, chunks, strict=True):
        _full_row(table, prompt)
        # The split already consumed each prompt once — the once the form prints it.
        body = _without(chunk, ())
        if body:
            _text_rows(table, body)
        else:
            _full_row(table, "")
    doc.add_paragraph()


def _effort(
    doc: DocxDocument,
    form: CupeForm,
    heading: str,
    text: str,
    rows: Sequence[WJQRow],
) -> None:
    """Section 10: the author's effort rows (type / duration / frequency, J2), padded
    to the form's blank rows, then any further text for the section."""
    instructions = form.instructions.get("effort", ())
    table = _table(doc, len(form.effort_columns))
    _heading_row(table, heading, instructions)
    _row(table, form.effort_columns, bold=True)
    _table_rows(
        table,
        [[r.item, r.duration or "", r.frequency or ""] for r in rows],
        form.blank_rows["effort"],
    )
    _text_rows(table, _without(text, instructions))
    doc.add_paragraph()


def _working_conditions(
    doc: DocxDocument, form: CupeForm, heading: str, text: str, extra: WJQForm
) -> None:
    """Section 11's three sub-tables. The author's rows (J2) fill each one — location /
    hours, condition / frequency, hazard / frequency. A stored section's text is split
    at the three prompts; each part is printed under its own sub-table (column headings
    kept — see :func:`_without`)."""
    parts = (
        (
            form.working_conditions_location_prompt,
            form.working_conditions_location_columns,
            "working_conditions_location",
            [[r.item, r.hours or ""] for r in extra.locations],
        ),
        (
            form.working_conditions_uncomfortable_prompt,
            form.working_conditions_uncomfortable_columns,
            "working_conditions_uncomfortable",
            [[r.item, r.frequency or ""] for r in extra.uncomfortable],
        ),
        (
            form.working_conditions_hazards_prompt,
            form.working_conditions_hazards_columns,
            "working_conditions_hazards",
            [[r.item, r.frequency or ""] for r in extra.hazards],
        ),
    )
    chunks = _split_at(text, [p for p, _, _, _ in parts])
    width = max(len(cols) for _, cols, _, _ in parts)
    table = _table(doc, width)
    _heading_row(table, heading, ())
    for (prompt, columns, rows_key, rows), chunk in zip(parts, chunks, strict=True):
        _full_row(table, prompt)
        _row(table, columns, bold=True)
        # The split already consumed each prompt once — the once the form prints it. A
        # source form that REPEATS a prompt (MEASURED: one prints the "uncomfortable"
        # prompt where the hazards prompt belongs) keeps the repeat.
        minimum = form.blank_rows[rows_key]  # type: ignore[index]
        if rows:
            # The author's rows sit under the column headings, as on HR's form.
            _table_rows(table, rows, minimum)
            _text_rows(table, _without(chunk, ()))
        else:
            # Unchanged from J1 for a role with no rows: its text, then blank rows.
            _text_rows(table, _without(chunk, ()))
            _table_rows(table, [], minimum)
    doc.add_paragraph()


def _qualifications(
    doc: DocxDocument, jd: SFUJobDescription, form: CupeForm, heading: str
) -> None:
    """Section 13, block by block as HR's form lays it out. Each qualification prints
    in the block ``cupe_form.qualification_blocks`` assigns its kind (HR-239)."""
    blocks: dict[str, list[str]] = {
        "education": [],
        "experience": [],
        "skills": [],
        "requirements": [],
    }
    for qual in jd.qualifications:
        if qual.text.strip():
            blocks[form.qualification_blocks[qual.kind]].append(_qual_text(qual))
    extra = jd.wjq_form or WJQForm()
    # A J2 draft keeps its Occupational Requirement(s) on the form block (Q9, option C);
    # they print after any `ability` / `security` lines HR-239 places there.
    blocks["requirements"] += [r for r in extra.occupational_requirements if r.strip()]

    table = _table(doc, 2)
    _heading_row(table, heading, form.instructions.get("qualifications", ()))
    _full_row(table, form.qualifications_education_prompt)
    _rows_full(table, blocks["education"])
    _full_row(table, form.qualifications_certification_prompt)
    _rows_full(table, list(extra.certification))
    _full_row(table, form.qualifications_experience_prompt)
    years = "" if extra.experience_years is None else str(extra.experience_years)
    _row(table, [years, form.qualifications_experience_unit])
    _rows_full(table, blocks["experience"])
    _full_row(table, form.qualifications_equivalency_heading, bold=True)
    # Option 0 = "is acceptable", option 1 = "is not acceptable because:" (HR's order).
    acceptable = extra.equivalency_acceptable
    chosen = -1 if acceptable is None else (0 if acceptable else 1)
    for i, option in enumerate(form.qualifications_equivalency_options):
        _full_row(table, f"{_box(form, i == chosen)} {option}")
    _full_row(table, extra.equivalency_rationale or "")
    doc.add_paragraph()

    for prompt, key in (
        (form.qualifications_skills_prompt, "skills"),
        (form.qualifications_requirements_prompt, "requirements"),
    ):
        box = _table(doc, 1)
        _full_row(box, prompt)
        for line in blocks[key]:
            doc.add_paragraph(line)
        if not blocks[key]:
            doc.add_paragraph()
        doc.add_paragraph()


def _rows_full(table: Table, lines: Sequence[str]) -> None:
    for line in lines:
        _full_row(table, line)
    if not lines:
        _full_row(table, "")


def _approval(doc: DocxDocument, form: CupeForm, heading: str) -> None:
    """Section 14 — ALWAYS blank lines under HR's labels (non-negotiable #1)."""
    table = _table(doc, 2)
    _full_row(table, heading, bold=True)
    for statement, labels in (
        (form.approval_supervisor_statement, form.approval_supervisor_labels),
        (form.approval_employee_statement, form.approval_employee_labels),
    ):
        _full_row(table, statement, bold=True)
        name, signature, *rest = (*labels, "", "")
        _row(table, ["", ""])
        _row(table, [name, signature])
        _row(table, ["", ""])
        _row(table, ["", rest[0]])


# ── the entry point ────────────────────────────────────────────────────────────


def render_cupe_docx(jd: SFUJobDescription, *, rules: Rules | None = None) -> bytes:
    """Compose ``jd`` on HR's CUPE "WJQ Custom — Part 1" form and return the bytes.

    Public so a caller CAN name the form, but callers should use
    :func:`~src.jd_export.render.render_sfu_docx`, which picks the form from the JD."""
    rulebook = rules if rules is not None else get_rules()
    form = rulebook.cupe_form
    wjq = rulebook.wjq
    headings = _headings(wjq)
    letters = {meaning: letter for letter, meaning in wjq.frequency_markers.items()}
    # The CUPE form block of an authored draft (J2); empty for a parsed role, which
    # then prints exactly as J1 did.
    extra = jd.wjq_form or WJQForm()

    context = _context_sections(jd.additional_context, wjq)
    # The model's fields that have no slot of their own on the CUPE form (HR-239).
    extras: dict[str, list[str]] = {}
    if jd.relationships is not None and (jd.relationships.supervisory or "").strip():
        extras.setdefault(form.supervisory_section, []).append(
            jd.relationships.supervisory or ""
        )
    extras.setdefault(form.decision_making_section, []).extend(jd.decision_making)
    extras.setdefault(form.problem_solving_section, []).extend(jd.problem_solving)

    def section_text(section: WjqSection) -> str:
        lines = [context.get(section, ""), *extras.get(section, [])]
        return "\n".join(line.strip() for line in lines if line.strip())

    doc = Document()
    font = doc.styles["Normal"].font
    font.name = form.font_name
    font.size = Pt(form.font_size_pt)

    for line in form.banner:
        _bold(doc.add_paragraph(), line)
    doc.add_paragraph()

    _identification(doc, jd, form, headings["position_identification"])
    doc.add_paragraph()

    summary = _table(doc, 1)
    _heading_row(
        summary,
        headings["position_summary"],
        form.instructions.get("position_summary", ()),
    )
    for line in (jd.position_summary or "").splitlines() or [""]:
        doc.add_paragraph(line.strip())
    doc.add_paragraph()

    majors = _table(doc, 1)
    _heading_row(
        majors,
        headings["major_functions"],
        form.instructions.get("major_functions", ()),
    )
    for number, duty in enumerate(jd.duties, start=1):
        doc.add_paragraph(f"{number}. {_duty_text(duty, letters)}")
    if not jd.duties:
        doc.add_paragraph()
    doc.add_paragraph()

    minors = _table(doc, 1)
    _heading_row(
        minors,
        headings["minor_functions"],
        form.instructions.get("minor_functions", ()),
    )
    # A J2 draft keeps its minor functions apart; a parsed role has none (the parser
    # merges sections 3 and 4, so its duties all printed under MAJOR FUNCTIONS above).
    for number, duty in enumerate(extra.minor_functions, start=1):
        doc.add_paragraph(f"{number}. {_duty_text(duty, letters)}")
    if not extra.minor_functions:
        doc.add_paragraph()
    doc.add_paragraph()

    for section in (
        "level_of_independence",
        "training_exercised",
        "direction_exercised",
    ):
        _checkbox_section(
            doc,
            headings[section],
            form.instructions.get(section, ()),
            form.checkbox_options[section],
            form,
            getattr(extra, section),
            section_text(section),
        )
    _contacts(doc, jd, form, headings["internal_external_contacts"])
    _impact_of_errors(
        doc, form, headings["impact_of_errors"], section_text("impact_of_errors")
    )
    _effort(doc, form, headings["effort"], section_text("effort"), extra.effort)
    _working_conditions(
        doc,
        form,
        headings["working_conditions"],
        section_text("working_conditions"),
        extra,
    )
    _checkbox_section(
        doc,
        headings["continuing_education"],
        form.instructions.get("continuing_education", ()),
        form.checkbox_options["continuing_education"],
        form,
        extra.continuing_education,
        section_text("continuing_education"),
    )
    _qualifications(doc, jd, form, headings["qualifications"])
    _approval(doc, form, headings["approval_review"])

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
